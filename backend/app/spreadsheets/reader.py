"""Reading uploaded spreadsheets: type sniffing, container guards and row iteration.

SECURITY (an upload is untrusted binary input):

* **The type comes from the bytes** (``sniff``), never from the Content-Type header: xlsx is a zip that
  holds ``[Content_Types].xml`` and ``xl/workbook.xml``; csv is text. The file name can only get a file
  refused (macro-enabled extensions), and it lets a delimiter-free text file named ``.csv`` count as a
  one-column csv; it never makes binary content acceptable.
* **Macro-enabled workbooks are refused:** any ``vbaProject.bin`` part, Excel 4.0 macro sheets, or a
  ``macroEnabled``/``vbaProject`` content type (xlsm, xltm, xlam). xlsb has no ``xl/workbook.xml``.
* **Zip bombs:** the central directory is checked before anything is inflated: at most
  ``MAX_ZIP_ENTRIES`` entries and ``MAX_UNCOMPRESSED_BYTES`` declared in total. CPython's ``zipfile``
  never returns more than an entry's declared size (a lying entry fails its CRC check), so the declared
  total bounds what openpyxl can inflate.
* **Nothing in a workbook is executed or evaluated:** openpyxl runs with ``read_only=True,
  data_only=True`` (cached formula results only, formulas are never computed), ``keep_links=False``,
  and VBA is never loaded. openpyxl parses XML with the standard library (expat, no external entities);
  expat 2.4.1 and newer refuse billion-laughs style entity expansion, and xlsx files are refused on an
  older expat (``_EXPAT_SAFE``).
* **Bounded work:** rows are read with explicit ``max_row``/``max_col``. Without them openpyxl pads every
  row to the sheet's declared dimension and back-fills row-number gaps one row at a time, and both
  numbers come from the file. At most ``MAX_SHEETS`` sheets are read and at most ``scan_limit(max_rows)``
  rows of each are scanned. ``csv.Sniffer`` sees at most ``SNIFF_SAMPLE_CHARS`` characters, because its
  quote detection is quadratic in the sample size and holds the GIL while it runs.
* **Damaged files fail soft:** anything openpyxl or ``csv`` raises while reading becomes
  ``UnsupportedFileType`` (logged with the exception type only, never content).

Values are normalised once, here, so the inspection describes exactly what ``iter_rows`` returns:
numbers written with Persian or Arabic-Indic digits (``"۱۲٬۵۰۰"``) become numbers, ``true``/``false``
become booleans, dates become ISO 8601 strings, empty cells become ``None``. Digit strings with a leading
zero (phone numbers, codes) or more than 15 digits (card numbers, ids) stay text.
"""

import csv
import io
import itertools
import logging
import math
import re
import zipfile
from collections.abc import Callable, Iterator, Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta, timezone
from typing import Any
from xml.parsers import expat

import openpyxl
from openpyxl.packaging.manifest import Manifest
from openpyxl.packaging.relationship import get_rels_path
from openpyxl.reader.workbook import WorkbookParser
from openpyxl.xml.constants import (
    ARC_CORE,
    ARC_CUSTOM,
    ARC_STYLE,
    SHARED_STRINGS,
    SHEET_MAIN_NS,
    XLSM,
    XLSX,
    XLTM,
    XLTX,
)
from openpyxl.xml.functions import fromstring

from app.botspec.records import normalize_digits
from app.spreadsheets.errors import (
    DAMAGED_MESSAGE,
    MACRO_MESSAGE,
    TOO_COMPLEX_MESSAGE,
    EmptyWorkbook,
    FileTooLarge,
    SpreadsheetError,
    UnsupportedFileType,
    too_many_columns_message,
)
from app.spreadsheets.storage import FileKind

log = logging.getLogger(__name__)

Cell = str | int | float | bool | None

DEFAULT_MAX_ROWS = 50_000
DEFAULT_MAX_COLUMNS = 100
MAX_ZIP_ENTRIES = 200
MAX_UNCOMPRESSED_BYTES = 60 * 1024 * 1024
MAX_CONTENT_TYPES_BYTES = 1024 * 1024
MAX_SHEETS = 10
EXCEL_MAX_ROWS = 1_048_576
SNIFF_SAMPLE_CHARS = 8192
CSV_DELIMITERS = ",;\t|"
MAX_LABEL_CHARS = 64
MAX_INT_DIGITS = 15  # longer digit strings are identifiers; also keeps integers exact in JavaScript
MAX_SAFE_INTEGER = 2**53 - 1

ZIP_MAGIC = b"PK\x03\x04"
CONTENT_TYPES_PART = "[Content_Types].xml"
WORKBOOK_PART = "xl/workbook.xml"
MACRO_EXTENSIONS = (".xlsm", ".xltm", ".xlam", ".xlsb")
CSV_SHEET_NAME = "csv"

# Structural limits for the XML inside an xlsx, enforced by ``check_xlsx_structure`` (see its docstring).
MAX_XML_DEPTH = 64
MAX_TREE_ELEMENTS = 150_000  # all parts openpyxl parses into whole trees, together
MAX_SUBTREE_ELEMENTS = 100_000  # one row, one shared string, one other element of a streamed part
MAX_RETAINED_ELEMENTS = 100_000  # elements of a streamed part that openpyxl never releases
MAX_SHARED_STRINGS = 500_000
_SCAN_CHUNK = 64 * 1024
_WORKBOOK_TYPES = (XLTM, XLTX, XLSM, XLSX)  # openpyxl's search order
_ROW_TAG = f"{SHEET_MAIN_NS} row"  # expat's form of the tags openpyxl releases while streaming
_SI_TAG = f"{SHEET_MAIN_NS} si"

_EXPAT_SAFE = expat.version_info >= (2, 4, 1)
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
# Control characters become spaces; invisible formatting marks (the Arabic letter mark, zero-width space,
# left-to-right/right-to-left marks, bidi embeddings, overrides and isolates, BOM) are removed. Used for
# labels and display file names; ZWNJ stays (Persian spelling).
INVISIBLE_CHARS: dict[int, str | None] = {
    **dict.fromkeys([*range(0x00, 0x20), *range(0x7F, 0xA0)], " "),
    **dict.fromkeys([0x061C, 0x200B, 0x200E, 0x200F, *range(0x202A, 0x202F), *range(0x2066, 0x206A), 0xFEFF]),
}
_NUMBER_RE = re.compile(r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?", re.ASCII)
_NUMBER_PUNCT = str.maketrans({"٫": ".", "٬": ",", "−": "-"})
_DATE_RE = re.compile(
    r"(\d{4})([-/])(\d{1,2})\2(\d{1,2})"
    r"(?:[T ](\d{1,2}):(\d{2})(?::(\d{2})(?:\.(\d{1,6}))?)?)?"
    r"(Z|[+-]\d{2}:?\d{2})?",
    re.ASCII,
)
_DATE_YEARS = range(1900, 2200)  # "1405/07/14" is a Jalali date: it stays text
_NUMERIC_START = frozenset("+-−0123456789۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩")


# --- Value normalisation -------------------------------------------------------------------------


def parse_number(text: str) -> int | float | None:
    """A number written in ``text`` (ASCII, Persian or Arabic-Indic digits, ``,``/``٬`` thousands
    separators in groups of three, ``.``/``٫`` decimal point), or None. A leading zero or more than
    ``MAX_INT_DIGITS`` integer digits means an identifier, not a number."""
    candidate = normalize_digits(text.strip()).translate(_NUMBER_PUNCT)
    if not _NUMBER_RE.fullmatch(candidate):
        return None
    candidate = candidate.replace(",", "")
    integer_part = candidate.lstrip("+-").partition(".")[0]
    if (len(integer_part) > 1 and integer_part[0] == "0") or len(integer_part) > MAX_INT_DIGITS:
        return None
    if "." not in candidate:
        return int(candidate)
    value = float(candidate)
    return value if math.isfinite(value) else None


def parse_iso_datetime(text: str) -> str | None:
    """``text`` as a canonical ISO 8601 string if it is a Gregorian date or date-time (``2026-10-05``,
    ``2026/10/05 14:30``, ``2026-10-05T14:30:00+03:30``; any digits), else None. A naive midnight is
    rendered as a plain date, as are date-only cells from xlsx files."""
    match = _DATE_RE.fullmatch(normalize_digits(text.strip()))
    if match is None:
        return None
    year, _, month, day, hour, minute, second, fraction, offset = match.groups()
    if int(year) not in _DATE_YEARS:
        return None
    try:
        day_value = date(int(year), int(month), int(day))
        if hour is None:
            return day_value.isoformat() if offset is None else None
        moment = datetime.combine(
            day_value,
            time(int(hour), int(minute), int(second or 0), int((fraction or "0").ljust(6, "0"))),
            tzinfo=_offset(offset),
        )
    except ValueError:
        return None
    return _iso(moment)


def _offset(text: str | None) -> timezone | None:
    if text is None:
        return None
    if text == "Z":
        return UTC
    digits = text[1:].replace(":", "")
    delta = timedelta(hours=int(digits[:2]), minutes=int(digits[2:]))
    if int(digits[2:]) >= 60:
        raise ValueError("invalid offset")
    return timezone(-delta if text[0] == "-" else delta)  # ValueError beyond +-24h


def _iso(moment: datetime) -> str:
    if moment.tzinfo is None and moment.time() == time(0):
        return moment.date().isoformat()
    return moment.isoformat()


def _duration(value: timedelta) -> str:
    seconds = int(value.total_seconds())
    sign, seconds = ("-", -seconds) if seconds < 0 else ("", seconds)
    return f"{sign}{seconds // 3600}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


def normalize_cell(value: Any) -> Cell:
    """One cell value as returned by ``iter_rows`` (see the module docstring)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value if abs(value) <= MAX_SAFE_INTEGER else str(value)
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, datetime):
        return _iso(value)
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, timedelta):
        return _duration(value)
    text = (value if isinstance(value, str) else str(value)).strip()
    if not text:
        return None
    first = text[0]
    if first in _NUMERIC_START:  # numbers and dates start with a digit or a sign
        number = parse_number(text)
        if number is not None:
            return number
        return parse_iso_datetime(text) or text
    if first in "tTfF":
        lowered = text.casefold()
        if lowered in ("true", "false"):
            return lowered == "true"
    return text


def display(value: Cell) -> str:
    """A normalised value as text (samples, labels)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return repr(value)
    return str(value)


def truncate(text: str, limit: int) -> str:
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _label(value: Any) -> str | None:
    """A header cell or sheet title as a display label: whitespace collapsed, control and bidi
    formatting characters removed, at most ``MAX_LABEL_CHARS``; None when nothing is left."""
    if value is None:
        return None
    text = value if isinstance(value, str) else display(normalize_cell(value))
    text = " ".join(text.translate(INVISIBLE_CHARS).split())
    return truncate(text, MAX_LABEL_CHARS) if text else None


def _unique(name: str, used: set[str]) -> str:
    candidate, n = name, 2
    while candidate in used:
        candidate = f"{name}_{n}"
        n += 1
    used.add(candidate)
    return candidate


def _headers(labels: Sequence[str | None]) -> list[str]:
    """Column names from the header row: up to its last non-empty cell; a blank cell is ``col_<n>``
    (1-based position) and a repeated name gets ``_2``, ``_3`` ..."""
    width = max((i + 1 for i, label in enumerate(labels) if label), default=0)
    used: set[str] = set()
    return [_unique(labels[i] or f"col_{i + 1}", used) for i in range(width)]


def _raw_empty(row: Sequence[Any]) -> bool:
    return row.count(None) + row.count("") == len(row)


def scan_limit(max_rows: int) -> int:
    """How many rows of one sheet are scanned at most (header search and blank rows included)."""
    return max_rows * 2 + 1000


# --- Row streams ---------------------------------------------------------------------------------


class RowStream:
    """The data rows of one sheet, read lazily: lists of normalised cells, exactly as wide as the
    headers; blank rows are skipped. ``row_limit_hit`` turns true when rows were left unread (more than
    ``max_rows`` data rows, or the sheet goes on past the scan limit); read it after the rows."""

    def __init__(self, raw: Iterator[Sequence[Any]], *, width: int, max_rows: int, scan_left: int) -> None:
        self.rows_read = 0
        self.row_limit_hit = False
        self._raw = raw
        self._closers: list[Callable[[], None]] = []
        self._released = False
        self._rows = self._generate(width, max_rows, scan_left)

    def __iter__(self) -> "RowStream":
        return self

    def __next__(self) -> list[Cell]:
        return next(self._rows)

    def on_close(self, closer: Callable[[], None]) -> None:
        self._closers.append(closer)

    def close(self) -> None:
        # Re-entrant: a closer may close this stream again from inside the generator's own cleanup.
        if self._released:
            return
        self._rows.close()
        self._release()

    def _release(self) -> None:
        if self._released:
            return
        self._released = True  # first, so a closer that calls close() returns at once
        close = getattr(self._raw, "close", None)
        if close is not None:
            close()
        for closer in self._closers:
            closer()

    def _generate(self, width: int, max_rows: int, scan_left: int) -> Iterator[list[Cell]]:
        try:
            for raw_row in self._raw:
                if scan_left <= 0:
                    self.row_limit_hit = True
                    return
                scan_left -= 1
                if _raw_empty(raw_row):
                    continue
                cells = [normalize_cell(value) for value in itertools.islice(raw_row, width)]
                if cells.count(None) == len(cells):
                    continue
                if self.rows_read >= max_rows:
                    self.row_limit_hit = True
                    return
                cells.extend([None] * (width - len(cells)))
                self.rows_read += 1
                yield cells
        finally:
            self._release()


@dataclass
class SheetRows:
    name: str
    headers: list[str]
    rows: RowStream


def _guarded(rows: Iterator[Any]) -> Iterator[Any]:
    """``rows``, with any error raised while producing them turned into ``UnsupportedFileType``."""
    try:
        while True:
            try:
                row = next(rows)
            except StopIteration:
                return
            except SpreadsheetError:
                raise
            except Exception as exc:
                log.info("spreadsheet rows could not be read (%s)", type(exc).__name__)
                raise UnsupportedFileType(DAMAGED_MESSAGE) from None
            yield row
    finally:
        close = getattr(rows, "close", None)
        if close is not None:
            close()


def _open_sheet(name: str, raw_rows: Iterator[Any], *, max_rows: int, max_cols: int) -> SheetRows | None:
    """Find the header (the first non-empty row) and return the sheet, or None if it has none."""
    raw = _guarded(raw_rows)
    limit, scanned = scan_limit(max_rows), 0
    for raw_row in raw:
        scanned += 1
        if scanned > limit:
            break
        if _raw_empty(raw_row):
            continue
        labels = [_label(value) for value in itertools.islice(raw_row, max_cols)]
        if any(labels):
            headers = _headers(labels)
            rows = RowStream(raw, width=len(headers), max_rows=max_rows, scan_left=limit - scanned)
            return SheetRows(name=name, headers=headers, rows=rows)
    raw.close()
    return None


# --- Containers ----------------------------------------------------------------------------------


def _is_macro_part(name: str) -> bool:
    lowered = name.replace("\\", "/").lower()
    return lowered.rsplit("/", 1)[-1] == "vbaproject.bin" or "macrosheets/" in lowered


def check_xlsx_container(data: bytes) -> None:
    """Refuse a zip that is not a plain (macro-free) xlsx workbook, or that would inflate past the
    limits, before anything in it is decompressed (beyond ``[Content_Types].xml``, capped at 1 MiB)."""
    if not _EXPAT_SAFE:
        log.error("xlsx upload refused: expat %s is older than 2.4.1", expat.version_info)
        raise UnsupportedFileType(DAMAGED_MESSAGE)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ZIP_ENTRIES:
                raise FileTooLarge()
            if sum(max(info.file_size, 0) for info in infos) > MAX_UNCOMPRESSED_BYTES:
                raise FileTooLarge()
            if any(_is_macro_part(info.filename) for info in infos):
                raise UnsupportedFileType(MACRO_MESSAGE)
            names = {info.filename for info in infos}
            if CONTENT_TYPES_PART not in names or WORKBOOK_PART not in names:
                raise UnsupportedFileType()
            if archive.getinfo(CONTENT_TYPES_PART).file_size > MAX_CONTENT_TYPES_BYTES:
                raise UnsupportedFileType(DAMAGED_MESSAGE)
            content_types = archive.read(CONTENT_TYPES_PART).lower()
    except SpreadsheetError:
        raise
    except Exception as exc:
        log.info("upload is not a readable zip container (%s)", type(exc).__name__)
        raise UnsupportedFileType(DAMAGED_MESSAGE) from None
    if b"macroenabled" in content_types or b"vbaproject" in content_types:
        raise UnsupportedFileType(MACRO_MESSAGE)


# --- XML structure (constant-memory pre-scan) ----------------------------------------------------


def _too_complex() -> FileTooLarge:
    return FileTooLarge(TOO_COMPLEX_MESSAGE)


def _refuse_doctype(*_: object) -> None:
    raise UnsupportedFileType(DAMAGED_MESSAGE)  # no OOXML part has a DTD; entities are never processed


class _PartScan:
    """Expat handlers that check one XML part without building a tree (constant memory).

    A tree part (``release`` None) adds its elements to the workbook's shared ``budget``. A streamed
    part is read by openpyxl with ``iterparse``, which empties each ``release`` element (a row of a
    worksheet, an ``si`` of the shared strings) once read and keeps everything else; so each released
    element may hold at most ``MAX_SUBTREE_ELEMENTS``, and at most ``MAX_RETAINED_ELEMENTS`` may exist
    outside them. Row numbers must increase (as the format requires): openpyxl stops reading a sheet
    at ``max_row`` only then, and otherwise keeps every emptied row element of the sheet."""

    def __init__(self, release: str | None, budget: list[int]) -> None:
        self.release = release
        self.budget = budget
        self.depth = self.count = self.retained = self.released = self.last_row = 0
        self.inside_from = 0  # element number where the open released element started; 0 if none

    def start(self, tag: str, attrs: dict[str, str]) -> None:
        self.count += 1
        self.depth += 1
        if self.depth > MAX_XML_DEPTH:
            raise _too_complex()
        if self.release is None:
            self.budget[0] += 1
            if self.budget[0] > MAX_TREE_ELEMENTS:
                raise _too_complex()
        elif self.inside_from:
            if tag == self.release or self.count - self.inside_from >= MAX_SUBTREE_ELEMENTS:
                raise _too_complex()
        elif tag == self.release:
            self.inside_from = self.count
            self.released += 1
            if self.release == _SI_TAG and self.released > MAX_SHARED_STRINGS:
                raise _too_complex()
            if self.release == _ROW_TAG:
                self._row_number(attrs.get("r"))
        else:
            self.retained += 1
            if self.retained > MAX_RETAINED_ELEMENTS:
                raise _too_complex()

    def end(self, tag: str) -> None:
        if self.inside_from and tag == self.release:
            self.inside_from = 0  # released elements never nest (refused above)
        self.depth -= 1

    def _row_number(self, value: str | None) -> None:
        if value is None:
            row = self.last_row + 1
        else:  # openpyxl's own reading of the attribute
            try:
                row = int(value)
            except ValueError:
                number = float(value)
                if not number.is_integer():
                    raise ValueError("invalid row number") from None
                row = int(number)
        if row <= self.last_row:
            raise UnsupportedFileType(DAMAGED_MESSAGE)
        self.last_row = row


def _scan_part(archive: zipfile.ZipFile, name: str, release: str | None, budget: list[int]) -> None:
    scan = _PartScan(release, budget)
    parser = expat.ParserCreate(namespace_separator=" ")
    parser.StartElementHandler = scan.start
    parser.EndElementHandler = scan.end
    parser.StartDoctypeDeclHandler = _refuse_doctype
    with archive.open(name) as stream:
        while chunk := stream.read(_SCAN_CHUNK):
            parser.Parse(chunk, False)
    parser.Parse(b"", True)


def _workbook_part_name(manifest: Manifest) -> str:
    """The part openpyxl loads as the workbook (``openpyxl.reader.excel._find_workbook_part``)."""
    for content_type in _WORKBOOK_TYPES:
        part = manifest.find(content_type)
        if part:
            return str(part.PartName)[1:]
    if {p.ContentType for p in manifest.Default} & set(_WORKBOOK_TYPES):
        return "xl/workbook.xml"
    raise UnsupportedFileType(DAMAGED_MESSAGE)


def check_xlsx_structure(archive: zipfile.ZipFile) -> None:
    """Refuse an xlsx whose XML would make openpyxl use unbounded memory, before openpyxl reads it.

    openpyxl parses most parts into whole element trees and keeps half-released trees while streaming
    worksheets and shared strings, so its memory grows 10 to 25 times faster than the inflated XML: a
    150 KB upload with one two-million-cell row took 1.4 GB, well inside ``MAX_UNCOMPRESSED_BYTES``.
    This streams the parts openpyxl will parse through expat (no tree) and enforces ``_PartScan``'s
    limits; DOCTYPEs are refused outright.

    Which part is what is decided the way openpyxl decides it, by openpyxl's own code: the manifest
    names the workbook and shared-strings parts, the workbook's ``<sheet>`` elements and relationships
    name the sheets. Those three small parts are bounded by the scan before openpyxl parses them. Tree
    parts: the manifest, the workbook and its relationships, styles, document properties, every
    sheet's relationships, chart sheets (and, when there are chart sheets, every other XML part, since
    their drawings and charts are parsed too). Streamed parts: the shared strings and the worksheets
    that ``iter_sheets`` reads. Parts openpyxl never opens in read-only mode (pivot caches, comments,
    calcChain ...) are not checked and never parsed."""
    names = set(archive.namelist())
    budget = [0]
    _scan_part(archive, CONTENT_TYPES_PART, None, budget)
    manifest = Manifest.from_tree(fromstring(archive.read(CONTENT_TYPES_PART)))
    workbook = _workbook_part_name(manifest)
    strings_part = manifest.find(SHARED_STRINGS)
    strings = str(strings_part.PartName)[1:] if strings_part else None
    rels = get_rels_path(workbook)
    if workbook not in names or (strings is not None and strings not in names):
        raise UnsupportedFileType(DAMAGED_MESSAGE)
    if rels in names:
        _scan_part(archive, rels, None, budget)
    _scan_part(archive, workbook, None, budget)

    parser = WorkbookParser(archive, workbook, keep_links=False)
    parser.parse()
    sheets = [(str(rel.Type or ""), str(rel.target)) for _, rel in parser.find_sheets()]
    worksheets = [target for kind, target in sheets if target in names and "chartsheet" not in kind]
    chartsheets = [target for kind, target in sheets if target in names and "chartsheet" in kind]

    tree_parts = {ARC_STYLE, ARC_CORE, ARC_CUSTOM, *chartsheets}
    tree_parts |= {get_rels_path(target) for target in worksheets + chartsheets}
    if chartsheets:
        streamed = {*worksheets, *([strings] if strings else [])}
        tree_parts |= {n for n in names if n.endswith((".xml", ".rels")) and n not in streamed}
    for part in sorted((tree_parts & names) - {CONTENT_TYPES_PART, rels, workbook}):
        _scan_part(archive, part, None, budget)
    if strings is not None:
        _scan_part(archive, strings, _SI_TAG, budget)
    for target in worksheets[:MAX_SHEETS]:
        _scan_part(archive, target, _ROW_TAG, budget)


def decode_text(data: bytes) -> str:
    """CSV bytes as text: UTF-8 (with or without BOM), else Windows-1256. Text holding control
    characters other than tab and line breaks is not a CSV file."""
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("cp1256")
        except UnicodeDecodeError:
            raise UnsupportedFileType() from None
    text = text.rstrip("\x1a")  # DOS end-of-file marker of old exports
    if _CONTROL_RE.search(text):
        raise UnsupportedFileType()
    return text


def _sniff_sample(text: str) -> str:
    if len(text) <= SNIFF_SAMPLE_CHARS:
        return text
    sample = text[:SNIFF_SAMPLE_CHARS]
    cut = sample.rfind("\n")
    return sample[: cut + 1] if cut > 0 else sample


def _fallback_delimiter(sample: str, *, single_column_ok: bool) -> str:
    """When ``csv.Sniffer`` finds no consistent delimiter (short files with a title line): the candidate
    present on the most lines if that is at least half of them; a one-column file when no line has any
    candidate and that is allowed."""
    lines = [line for line in sample.splitlines() if line.strip()]
    counts = {d: sum(d in line for line in lines) for d in CSV_DELIMITERS}
    best = max(CSV_DELIMITERS, key=lambda d: counts[d])  # ties: the first in CSV_DELIMITERS
    if counts[best] and counts[best] * 2 >= len(lines):
        return best
    if not counts[best] and single_column_ok:
        return ","
    raise UnsupportedFileType()


def csv_format(text: str, *, max_cols: int, single_column_ok: bool) -> tuple[str, bool]:
    """(delimiter, skipinitialspace) of CSV ``text``; ``UnsupportedFileType`` if it does not look like
    a table of at most ``max_cols`` columns, ``EmptyWorkbook`` if it has no non-blank line."""
    sample = _sniff_sample(text)
    if not sample.strip():
        raise EmptyWorkbook()
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=CSV_DELIMITERS)
        delimiter, skip_space = dialect.delimiter, bool(dialect.skipinitialspace)
    except csv.Error:
        delimiter, skip_space = _fallback_delimiter(sample, single_column_ok=single_column_ok), False
    reader = csv.reader(io.StringIO(sample, newline=""), delimiter=delimiter, skipinitialspace=skip_space)
    try:
        header = next((row for row in reader if any(cell.strip() for cell in row)), None)
    except csv.Error:
        raise UnsupportedFileType() from None
    if header is None:
        raise EmptyWorkbook()
    if max((i + 1 for i, cell in enumerate(header) if cell.strip()), default=0) > max_cols:
        raise UnsupportedFileType(too_many_columns_message(max_cols))
    return delimiter, skip_space


def sniff(data: bytes, filename: str, *, max_cols: int = DEFAULT_MAX_COLUMNS) -> FileKind:
    """The file kind judged from the bytes (see the module docstring), or a ``SpreadsheetError``."""
    name = (filename or "").strip().lower()
    if name.endswith(MACRO_EXTENSIONS):
        raise UnsupportedFileType(MACRO_MESSAGE)
    if data.startswith(ZIP_MAGIC):
        check_xlsx_container(data)
        return "xlsx"
    csv_format(decode_text(data), max_cols=max_cols, single_column_ok=name.endswith(".csv"))
    return "csv"


# --- Sheets and rows -----------------------------------------------------------------------------


def _open_workbook(data: bytes) -> Any:
    check_xlsx_container(data)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            check_xlsx_structure(archive)
        return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except SpreadsheetError:
        raise
    except Exception as exc:
        log.info("workbook could not be opened (%s)", type(exc).__name__)
        raise UnsupportedFileType(DAMAGED_MESSAGE) from None


def iter_sheets(
    data: bytes, kind: FileKind, *, max_rows: int = DEFAULT_MAX_ROWS, max_cols: int = DEFAULT_MAX_COLUMNS
) -> Iterator[SheetRows]:
    """Every sheet that has a header row, in workbook order: the first ``MAX_SHEETS`` worksheets of an
    xlsx file (chart sheets are not worksheets), or one sheet named ``"csv"``. Names are labels made
    unique. Read a sheet's rows before advancing to the next sheet."""
    if kind == "csv":
        text = decode_text(data)
        delimiter, skip_space = csv_format(text, max_cols=max_cols, single_column_ok=True)
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, skipinitialspace=skip_space)
        sheet = _open_sheet(CSV_SHEET_NAME, reader, max_rows=max_rows, max_cols=max_cols)
        if sheet is not None:
            try:
                yield sheet
            finally:
                sheet.rows.close()
        return
    if kind != "xlsx":
        raise ValueError(f"unknown file kind {kind!r}")
    workbook = _open_workbook(data)
    try:
        used: set[str] = set()
        max_row = min(scan_limit(max_rows), EXCEL_MAX_ROWS) + 1  # one past: tells whether rows remain
        for worksheet in workbook.worksheets[:MAX_SHEETS]:
            name = _unique(_label(worksheet.title) or "sheet", used)
            raw = worksheet.iter_rows(
                min_row=1, max_row=max_row, min_col=1, max_col=max_cols, values_only=True
            )
            sheet = _open_sheet(name, raw, max_rows=max_rows, max_cols=max_cols)
            if sheet is None:
                continue
            try:
                yield sheet
            finally:
                sheet.rows.close()
    finally:
        workbook.close()


def iter_rows(
    data: bytes,
    kind: FileKind,
    sheet: str | None = None,
    *,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_cols: int = DEFAULT_MAX_COLUMNS,
) -> tuple[list[str], RowStream]:
    """(headers, rows) of sheet ``sheet`` (by its name in the inspection; default: the first sheet with
    a header). ``rows.row_limit_hit`` tells whether rows were left unread. The workbook stays open until
    the rows are exhausted or ``rows.close()`` is called."""
    sheets = iter_sheets(data, kind, max_rows=max_rows, max_cols=max_cols)
    for found in sheets:
        if sheet is None or found.name == sheet:
            found.rows.on_close(sheets.close)
            return found.headers, found.rows
    if sheet is None:
        raise EmptyWorkbook()
    raise EmptyWorkbook(f"برگهٔ «{truncate(sheet, MAX_LABEL_CHARS)}» در فایل پیدا نشد.")
